"""生成示例数据与模板文件。

运行:python tools/make_samples.py
输出:samples/ 下的标准模板与外部导入示例表
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from ghd import excel_io

SAMPLES = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "samples")


def main() -> None:
    os.makedirs(SAMPLES, exist_ok=True)
    names = [
        "--28届--", "陈志远", "林晓峰", "王梓涵",
        "--29届--", "张三", "李四", "王五", "赵六", "周七", "张三",
        "--30届--", "刘一鸣", "孙浩然", "钱雨桐", "吴思远",
    ]
    # 标准总积分模板(含届别分隔行与一名同名"张三")
    excel_io.create_main_template(os.path.join(SAMPLES, "总积分模板_示例.xlsx"), names)
    # 外部导入示例:升降旗考勤(姓名/总次数)
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["姓名", "总次数"])
    for n, c in [("张三", 6), ("李四", 4), ("王五", 9), ("赵六", None), ("刘一鸣", 7)]:
        ws.append([n, c])
    wb.save(os.path.join(SAMPLES, "升降旗考勤示例.xlsx"))
    # 外部导入示例:日常训练考勤
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["姓名", "总次数"])
    for n, c in [("张三", 3), ("王五", 5), ("周七", 2), ("孙浩然", 4)]:
        ws.append([n, c])
    wb.save(os.path.join(SAMPLES, "日常训练考勤示例.xlsx"))
    # 外部导入示例:后勤部(智能夹心,中间明细列)
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["姓名", "搬水次数", "器材管理", "总积分"])
    for n, a, b, t in [("张三", 2, 1, 4.5), ("李四", 1, 0, 2.0), ("赵六", 3, 2, 6.5)]:
        ws.append([n, a, b, t])
    wb.save(os.path.join(SAMPLES, "部门积分示例_后勤.xlsx"))
    # 外部导入示例:综办部(无中间明细列)
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["姓名", "总积分"])
    for n, t in [("周七", 3.0), ("王五", 1.5), ("刘一鸣", 2.5)]:
        ws.append([n, t])
    wb.save(os.path.join(SAMPLES, "部门积分示例_综办.xlsx"))
    print("示例文件已生成到 samples/:")
    for f in sorted(os.listdir(SAMPLES)):
        print(" -", f)


if __name__ == "__main__":
    main()

"""核心层端到端功能验证:项目/锁/结算/导入/修改/撤销/备份/花名册。

运行:python tests/test_core.py
"""

import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from ghd import backups, excel_io, history_ops, import_parsers, project, recent
from ghd.history_ops import ProjectContext
from ghd.name_parser import parse_names
from ghd.roster import parse_roster
from ghd.scoring import EventForm, RoundCard

TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_tmp_test")
PASS, FAIL = 0, 0


def check(cond, label: str) -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [PASS] {label}")
    else:
        FAIL += 1
        print(f"  [FAIL] {label}")


def cell(path, row, col):
    wb = openpyxl.load_workbook(path)
    v = wb.active.cell(row=row, column=col).value
    wb.close()
    return v


def main() -> None:
    if os.path.isdir(TMP):
        shutil.rmtree(TMP, ignore_errors=True)
    os.makedirs(TMP, exist_ok=True)

    # ------------------------------------------------------------------
    print("== 1. 名单分隔符清洗器 ==")
    check(parse_names("张三，李四、王五, 赵六　钱七\n孙八") ==
          ["张三", "李四", "王五", "赵六", "钱七", "孙八"], "6 种分隔符解析")
    check(parse_names("  张三   张三  ") == ["张三"], "去重与去空")
    check(parse_names("") == [], "空串安全")

    # ------------------------------------------------------------------
    print("== 2. 总积分模板与花名册解析(届别状态机/同名检测) ==")
    main_xlsx = os.path.join(TMP, "总积分.xlsx")
    names = ["--30届--", "张三", "李四", "王五", "--29届--", "赵六", "周七", "张三"]
    excel_io.create_main_template(main_xlsx, names)
    matrix = excel_io.read_matrix(main_xlsx)
    roster = parse_roster(matrix[0], matrix[1:])
    check([m.name for m in roster.members] ==
          ["张三", "李四", "王五", "赵六", "周七", "张三"], "界别行不入人员")
    check(roster.members[0].generation == 30 and roster.members[3].generation == 29,
          "隐式打标 generation")
    check("张三" in roster.duplicate_names, "同名检测")
    check(cell(main_xlsx, 1, 6) == "汇总", "F 列为汇总")
    check(cell(main_xlsx, 1, 4) == "后勤部" and cell(main_xlsx, 1, 5) == "综办部",
          "D/E 列契约")

    # ------------------------------------------------------------------
    print("== 3. 工程配置 + 占用锁 ==")
    cfg_path = project.config_path_for(TMP, "2026秋季集训积分")
    cfg = project.default_config("2026秋季集训积分", main_xlsx)
    project.save_config(cfg_path, cfg)
    cfg = project.load_config(cfg_path)
    check(cfg["meta"]["project_name"] == "2026秋季集训积分", "工程配置持久化")
    ok, msg = project.acquire_lock(main_xlsx)
    check(ok, "建立占用锁")
    # 智能锁语义:同进程重复加锁视为自己残留,自动清理后成功
    ok2, msg2 = project.acquire_lock(main_xlsx)
    check(ok2 and "自动清理" in msg2, "同进程二次加锁自动清理")
    # 异进程活锁:真实阻断(信息含 PID)
    with open(project.excel_lock_path(main_xlsx), "w", encoding="utf-8") as f:
        f.write('{"pid": 12345, "time": "2026-10-02 12:00:00", "host": "X"}')
    from ghd import project as _proj
    _orig_alive = _proj.is_pid_alive
    _proj.is_pid_alive = lambda pid: True     # 模拟持锁进程存活
    try:
        ok3, msg3 = project.acquire_lock(main_xlsx)
    finally:
        _proj.is_pid_alive = _orig_alive
    check(not ok3 and "占用" in msg3 and "12345" in msg3, "异进程活锁被阻断(含 PID)")
    # 死进程残留:自动清理并重建
    with open(project.excel_lock_path(main_xlsx), "w", encoding="utf-8") as f:
        f.write('{"pid": 999999999, "time": "2020-01-01", "host": "X"}')
    ok4, msg4 = project.acquire_lock(main_xlsx)
    check(ok4 and "残留" in msg4, "死进程残留锁自动清理")
    project.release_lock(main_xlsx)
    check(not os.path.exists(project.excel_lock_path(main_xlsx)), "释放锁")

    ctx = ProjectContext(cfg_path, cfg)
    ok3, _ = project.acquire_lock(main_xlsx)

    # ------------------------------------------------------------------
    print("== 4. 活动积分结算引擎(计算法则) ==")
    form = EventForm(
        name="国庆出旗",
        chief_instructors=["张三", "周七"],
        training_cards=[RoundCard("training", 1.0, ["张三"], ["李四", "王五"])],
        rehearsal_cards=[RoundCard("rehearsal", 0.5, ["张三"], ["李四"])],
        formal_duration=2.0,
        formal_attendees=["张三", "王五"],
        field_duty=["赵六"],
    )
    errs = form.validate()
    check(not errs, f"表单校验通过 {errs if errs else ''}")
    scores = form.compute()
    expect = {"张三": 2 + 1.0 + 0.5 + 2.0, "周七": 2.0, "李四": 1.0 + 0.5,
              "王五": 1.0 + 2.0, "赵六": 0.5}
    check(all(abs(scores[k] - v) < 1e-9 for k, v in expect.items()) and
          set(scores) == set(expect), f"积分计算 {scores}")
    form2 = EventForm(name="X", chief_instructors=[])
    check(any("主教官" in e for e in form2.validate()), "主教官必填校验")
    form3 = EventForm(name="X", chief_instructors=["张"],
                      rehearsal_cards=[])
    check(any("彩排" in e for e in form3.validate()), "彩排至少 1 张校验")

    tx1 = history_ops.commit_event(ctx, "国庆出旗", scores)
    check(cell(main_xlsx, 1, 7) == "国庆出旗", "G 列追加活动列")
    r_zhang = roster.by_name["张三"][0].row
    check(abs(cell(main_xlsx, r_zhang, 7) - 5.5) < 1e-9, "单人分值填入新列")
    check(abs(cell(main_xlsx, r_zhang, 6) - 5.5) < 1e-9, "F 列汇总同步")
    check(tx1["status"] == "active" and tx1["is_first"], "事务日志(栈底/活跃)")

    # ------------------------------------------------------------------
    print("== 5. 升降旗考勤导入(倍率/空值清洗) ==")
    flag_xlsx = os.path.join(TMP, "升降旗考勤.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "总次数"])
    ws.append(["李四", 4])        # 日常 0.5 -> +2.0
    ws.append([None, 9])          # 姓名为空 -> 过滤
    ws.append(["赵六", None])     # 次数空 -> 0
    ws.append(["陌生人", 2])      # 不在主库
    wb.save(flag_xlsx)
    rows = import_parsers.parse_count_sheet(flag_xlsx)
    check(("李四", 4.0) in rows and ("赵六", 0.0) in rows and
          not any(n in (None, "") for n, _ in rows), "契约清洗(空行过滤/NaN->0)")
    deltas = {"李四": 2.0, "赵六": 0.0}
    tx2 = history_ops.commit_batch(ctx, "flag_duty", "升降旗考勤导入", deltas,
                                   target_col=import_parsers.target_col_for_scenes("flag_duty"),
                                   file="升降旗考勤.xlsx", rate=0.5)
    r_li = roster.by_name["李四"][0].row
    r_zhao = roster.by_name["赵六"][0].row
    check(abs(cell(main_xlsx, r_li, 3) - 2.0) < 1e-9, "C 列增量累加")
    check(abs(cell(main_xlsx, r_li, 6) - 3.5) < 1e-9, "F 列同步(1.5+2.0)")

    # ------------------------------------------------------------------
    print("== 6. 日常训练考勤导入 ==")
    dt_xlsx = os.path.join(TMP, "日常训练考勤模板.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "总次数"])
    ws.append(["王五", 3])
    wb.save(dt_xlsx)
    rows = import_parsers.parse_count_sheet(dt_xlsx)
    check(rows == [("王五", 3.0)], "日常训练契约解析")
    tx3 = history_ops.commit_batch(ctx, "daily_training", "日常训练考勤导入",
                                   {"王五": 3.0},
                                   target_col=import_parsers.target_col_for_scenes("daily_training"),
                                   file="日常训练考勤模板.xlsx", unit_score=1.0)
    r_wang = roster.by_name["王五"][0].row
    check(abs(cell(main_xlsx, r_wang, 2) - 3.0) < 1e-9, "B 列增量累加")
    check(abs(cell(main_xlsx, r_wang, 6) - 6.0) < 1e-9, "F 列同步(3.0+3.0)")

    # ------------------------------------------------------------------
    print("== 7. 部门积分导入(智能夹心 + 单次互斥锁) ==")
    dep_xlsx = os.path.join(TMP, "部门积分模板.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "搬水", "贴海报", "总积分"])   # 中间 2 列明细被忽略
    ws.append(["赵六", 1, 2, 3.5])
    wb.save(dep_xlsx)
    rows = import_parsers.parse_department_sheet(dep_xlsx)
    check(rows == [("赵六", 3.5)], "智能夹心提取(忽略中间列)")
    tx4 = history_ops.commit_batch(ctx, "department", "后勤部积分导入", {"赵六": 3.5},
                                   target_col=import_parsers.target_col_for_scenes("logistics"),
                                   file="部门积分模板.xlsx", department="logistics",
                                   stack_key="dept:logistics")
    r_zhao = roster.by_name["赵六"][0].row
    check(abs(cell(main_xlsx, r_zhao, 4) - 3.5) < 1e-9, "D 列累加")
    check(ctx.config["state"]["is_logistics_imported"] is True, "后勤互斥锁置位")
    check(abs(cell(main_xlsx, r_zhao, 6) - 4.0) < 1e-9, "F 列同步(0.5+3.5)")

    # ------------------------------------------------------------------
    print("== 8. 线性版本栈与部分修改(Delta Update) ==")
    hist = ctx.config["history"]
    check(len(hist) == 4 and all(t["status"] == "active" for t in hist), "各栈顶活跃")
    new_scores = dict(scores)
    new_scores["张三"] = 6.0     # +0.5
    new_scores["李四"] = 1.0     # -0.5
    tx5 = history_ops.commit_edit(ctx, tx1, new_scores)
    check(abs(cell(main_xlsx, r_zhang, 7) - 6.0) < 1e-9, "活动列差值累加")
    check(abs(cell(main_xlsx, r_zhang, 6) - 6.0) < 1e-9, "F 列差值同步(5.5+0.5)")
    check(abs(cell(main_xlsx, r_li, 7) - 1.0) < 1e-9, "李四差值 -0.5")
    stack = [t for t in ctx.config["history"] if t["stack_key"] == "activity:国庆出旗"]
    check(len(stack) == 2 and stack[0]["status"] == "locked" and
          stack[1]["status"] == "active", "原记录锁定,修改记录为栈顶")

    print("== 9. 撤销修改记录(反向差值) ==")
    history_ops.undo_tx(ctx, tx5)
    check(abs(cell(main_xlsx, r_zhang, 7) - 5.5) < 1e-9, "活动列恢复上一版本")
    check(abs(cell(main_xlsx, r_zhang, 6) - 5.5) < 1e-9, "F 列恢复")
    stack = [t for t in ctx.config["history"] if t["stack_key"] == "activity:国庆出旗"]
    check(stack[1]["status"] == "undone" and stack[0]["status"] == "active",
          "修改记录已撤销,原记录解锁")

    print("== 10. 撤销初次创建(扣 F + 物理删列) ==")
    history_ops.undo_tx(ctx, tx1)
    check(cell(main_xlsx, 1, 7) in (None, ""), "活动列已物理删除")
    check(abs(cell(main_xlsx, r_zhang, 6) - 0.0) < 1e-9, "F 列精准扣除")
    stack = [t for t in ctx.config["history"] if t["stack_key"] == "activity:国庆出旗"]
    check(stack[0]["status"] == "undone", "初次创建记录已撤销")

    print("== 11. 撤销部门批次(反向扣减 + 锁释放) ==")
    # 注:撤销活动批次时赵六的事件分 +0.5 已随 tx1 精准扣除,此时 F = 3.5
    history_ops.undo_tx(ctx, tx4)
    check(abs(cell(main_xlsx, r_zhao, 4) - 0.0) < 1e-9, "D 列反向扣除")
    check(abs(cell(main_xlsx, r_zhao, 6) - 0.0) < 1e-9, "F 列恢复(3.5-3.5)")
    check(ctx.config["state"]["is_logistics_imported"] is False, "部门锁重新点亮")

    print("== 12. 物理行序保持(排名只读) ==")
    col_a = [cell(main_xlsx, r, 1) for r in range(1, 10)]
    check(col_a == ["姓名"] + names, "原 Excel 物理行排版未被改动")

    print("== 13. 滚动备份机制 ==")
    bl = backups.list_backups(main_xlsx)
    check(len(bl) >= 5, f"备份已生成 {len(bl)} 份")
    check(all(os.path.isfile(b) for b in bl[:3]), "备份文件有效")
    for i in range(25):     # 触发滚动清理
        backups.make_backup(main_xlsx)
    bl = backups.list_backups(main_xlsx)
    check(len(bl) <= 20, f"滚动清理只保留最新 20 份(实际 {len(bl)})")

    print("== 14. 历史项目列表 ==")
    recent.add_recent(cfg_path, "2026秋季集训积分", main_xlsx)
    items = [it for it in recent.list_recent()
             if os.path.abspath(it["config_path"]) == os.path.abspath(cfg_path)]
    check(len(items) == 1 and items[0]["exists"], "登记与展示")
    recent.remove_recent(cfg_path)
    items = [it for it in recent.list_recent()
             if os.path.abspath(it["config_path"]) == os.path.abspath(cfg_path)]
    check(len(items) == 0, "清除记录")

    print("== 15. 界别分隔行格式兼容(第30届/--30届--/30届) ==")
    from ghd.roster import parse_roster as _pr
    from ghd.constants import GENERATION_RE as _GRE
    for label, gen in [("--30届--", 30), ("-- 30届 --", 30), ("--30届", 30),
                       ("30届--", 30), ("第30届", 30), ("第 30 届", 30), ("30届", 30)]:
        check(bool(_GRE.match(label)) and
              int(_GRE.match(label).group(1)) == gen, f"界别行「{label}」识别为 {gen}届")
    check(not _GRE.match("张三"), "人名不误判为界别行")
    m2 = excel_io.read_matrix(main_xlsx, max_col=2)
    r2 = _pr(m2[0], m2[1:])
    check(r2.members[0].generation == 30 and r2.members[3].generation == 29,
          "既有模板届别状态机不受影响")

    # ------------------------------------------------------------------
    print("== 16. 行号契约与汇总自动重算(回归 #3 错位写入) ==")
    xlsx2 = os.path.join(TMP, "行号契约.xlsx")
    names2 = ["--31届--", "甲一", "乙二", "--30届--", "丙三"]
    excel_io.create_main_template(xlsx2, names2)
    matrix2 = excel_io.read_matrix(xlsx2)
    roster2 = parse_roster(matrix2[0], matrix2[1:])
    # 行1=表头,行2=--31届--,行3=甲一,行4=乙二,行5=--30届--,行6=丙三
    check([m.row for m in roster2.members] == [3, 4, 6], "成员行号 = Excel 物理行")
    check(all(cell(xlsx2, m.row, 1) == m.name for m in roster2.members),
          "行号与姓名一一对应(不再错位到上一行)")

    cfg2_path = project.config_path_for(TMP, "行号契约")
    cfg2 = project.default_config("行号契约", xlsx2)
    project.save_config(cfg2_path, cfg2)
    cfg2 = project.load_config(cfg2_path)
    ctx2 = ProjectContext(cfg2_path, cfg2)
    project.acquire_lock(xlsx2)

    tx_e2 = history_ops.commit_event(ctx2, "运动会",
                                     {"甲一": 2.0, "乙二": 1.5, "丙三": 3.0})
    tx_f2 = history_ops.commit_batch(ctx2, "flag_duty", "升降旗导入",
                                     {"乙二": 2.0, "丙三": 1.0}, target_col=3,
                                     file="x.xlsx", rate=0.5)
    tx_d2 = history_ops.commit_batch(ctx2, "daily_training", "日常训练导入",
                                     {"甲一": 3.0}, target_col=2,
                                     file="x.xlsx", unit_score=1.0)
    # 汇总(F)自动重算 = B~E + G+ 之和,逐行验证
    def row_sum(path, row):
        s = sum((cell(path, row, c) or 0) for c in range(2, 6))       # B~E
        s += sum((cell(path, row, c) or 0) for c in range(7, 9))      # G+
        return s

    for m in roster2.members:
        f = cell(xlsx2, m.row, 6)
        s = row_sum(xlsx2, m.row)
        check(abs(f - s) < 1e-9, f"汇总重算一致:{m.name} F={f} = 各列和 {s}")
    check(abs(cell(xlsx2, roster2.members[0].row, 6) - 5.0) < 1e-9, "甲一汇总=活动2.0+日常3.0")
    check(abs(cell(xlsx2, roster2.members[1].row, 6) - 3.5) < 1e-9, "乙二汇总=活动1.5+升降旗2.0")
    check(abs(cell(xlsx2, roster2.members[2].row, 6) - 4.0) < 1e-9, "丙三汇总=活动3.0+升降旗1.0")
    # 撤销活动批次(物理删列)后,汇总仍 = 各列之和
    history_ops.undo_tx(ctx2, tx_e2)
    check(cell(xlsx2, 1, 7) in (None, ""), "活动列已删除")
    for m in roster2.members:
        f = cell(xlsx2, m.row, 6)
        s = row_sum(xlsx2, m.row)
        check(abs(f - s) < 1e-9, f"撤销后汇总仍一致:{m.name} F={f}")
    check(abs(cell(xlsx2, roster2.members[0].row, 6) - 3.0) < 1e-9, "甲一撤销后=日常3.0")
    project.release_lock(xlsx2)

    project.release_lock(main_xlsx)
    print(f"\n==== 结果:PASS {PASS} / FAIL {FAIL} ====")
    if FAIL:
        sys.exit(1)


if __name__ == "__main__":
    main()
